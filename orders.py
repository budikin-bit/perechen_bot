# -*- coding: utf-8 -*-
"""
Чтение приказов в формате RTF или DOCX → таблица строк + реквизиты приказа.

    from orders import read_order
    order = read_order("приказ.rtf")      # или .docx
    order.meta   # {'ministry': 'Минобрнауки России', 'number': '521', 'date': '31.08.2026', …}
    order.rows   # [[ячейка1, …, ячейка6], …] — все строки таблиц с ≥ 6 колонок

RTF читается собственным читалкой (без LibreOffice и без сторонних пакетов),
поэтому сборка каталога работает и на Windows. DOCX читается через python-docx.

Особенность вертикально объединённых ячеек:
  • в RTF продолжение объединённой ячейки — пустая строка;
  • python-docx повторяет в них текст первой ячейки.
Парсеры (parse_minobr.py, parse_minpros.py) рассчитаны на оба варианта.
"""
from __future__ import annotations

import codecs
import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Tuple

# ───────────────────────── кавычки ─────────────────────────

_OPEN_AFTER = set(" \t\n(«[—–-/;:")


def fix_quotes(s: str) -> str:
    """
    Прямые кавычки " → «ёлочки» (с учётом вложенности: «…«…»…»).
    В официальных текстах кавычки прямые; без этого не работают разбор
    организаторов и аббревиатуры.
    """
    if '"' not in s:
        return s
    out: List[str] = []
    for i, ch in enumerate(s):
        if ch != '"':
            out.append(ch)
            continue
        prev = out[-1] if out else ""
        nxt = s[i + 1] if i + 1 < len(s) else ""
        opening = (not prev or prev in _OPEN_AFTER) and nxt and not nxt.isspace()
        out.append("«" if opening else "»")
    res = "".join(out)
    return res.replace("“", "«").replace("”", "»").replace("„", "«")


def clean_cell(s: str) -> str:
    """Пробелы нормализуются построчно: переводы строк внутри ячейки сохраняются."""
    s = re.sub(r"[\udc80-\udcff]", "", s).replace("\xa0", " ").replace("\r", "")
    lines = [re.sub(r"[ \t]+", " ", ln).strip() for ln in s.split("\n")]
    return fix_quotes("\n".join(ln for ln in lines if ln))


# ───────────────────────── RTF ─────────────────────────

_SKIP_DEST = {
    "fonttbl", "colortbl", "stylesheet", "info", "pict", "shp", "shpinst", "shptxt",
    "shprslt", "nonshppict", "header", "footer", "headerl", "headerr", "headerf",
    "footerl", "footerr", "footerf", "listtable", "listoverridetable", "rsidtbl",
    "generator", "themedata", "colorschememapping", "latentstyles", "datastore",
    "xmlnstbl", "fldinst", "bkmkstart", "bkmkend", "pgptbl", "revtbl", "filetbl", "private",
}
_SYMBOLS = {"emdash": "—", "endash": "–", "lquote": "‘", "rquote": "’",
            "ldblquote": "«", "rdblquote": "»", "bullet": "•", "enspace": " ", "emspace": " "}
_TOK = re.compile(
    r"\\'([0-9a-fA-F]{2})|\\([a-zA-Z]+)(-?\d+)? ?|\\(.)|([{}])|([^\\{}]+)", re.S)
_CPG_RE = re.compile(r"\\ansicpg(\d+)")


_ENC = "cp1251"


def _codepage(head) -> str:
    """Кодировка из \\ansicpgN (по умолчанию и для неизвестных — cp1251)."""
    if isinstance(head, bytes):
        head = head.decode("latin-1")
    m = _CPG_RE.search(head[:4000])
    enc = f"cp{m.group(1)}" if m else "cp1251"
    try:
        codecs.lookup(enc)
    except LookupError:
        enc = "cp1251"
    return enc


def _decode_rtf(raw: bytes) -> str:
    """Текст RTF; байты без символа в кодировке сохраняются (surrogateescape) —
    это нужно, чтобы посчитать хэш встроенных картинок по исходным байтам."""
    global _ENC
    _ENC = _codepage(raw[:4000])
    return raw.decode(_ENC, "surrogateescape")


def _iter_tokens(s: str):
    """Токены RTF. Двоичные данные после \\binN пропускаются и отдаются целиком."""
    pos, n = 0, len(s)
    while pos < n:
        m = _TOK.match(s, pos)
        if not m:
            pos += 1
            continue
        pos = m.end()
        hexv, word, num, sym, brace, text = m.groups()
        if word == "bin" and num:
            k = int(num)
            data = s[pos:pos + k]
            pos += k
            yield (None, "bin", data, None, None, None)
        else:
            yield hexv, word, num, sym, brace, text


_HEX_DIGITS = re.compile(r"[^0-9a-fA-F]")
_STRUCT = {"par", "line", "tab", "cell", "row", "nestcell", "nestrow", "trowd"}


def _rtf_events(s: str):
    """
    Разбор RTF в поток событий (общий для таблиц и реквизитов):
      ("text", str) · ("img", sha1[:12]) · ("par", в_таблице) · ("line",) · ("tab",)
      ("cell",) · ("row",) · ("nestcell",) · ("nestrow",) · ("trowd", в_таблице) · ("end", в_таблице)

    Unicode: \\ucN действует в пределах группы (по умолчанию 1); после \\uN
    пропускаются N символов замены — обычный текст, \\'xx (один символ) или
    управляющий символ; управляющее слово или скобка группы пропуск прекращает.
    Отрицательные \\uN → +65536; пара суррогатов UTF-16 склеивается в один символ,
    одиночный суррогат → U+FFFD. Байты \\'xx декодируются кодировкой \\ansicpgN
    (последовательность байтов — целиком, поэтому работают и двухбайтовые кодировки).

    Картинки: \\pict с \\binN — хэш двоичных данных; \\pict с шестнадцатеричными
    данными — хэш декодированных байтов. Картинки внутри пропускаемых групп
    (\\nonshppict и т. п.) не учитываются; {\\*\\shppict …} — учитывается.
    """
    enc = _codepage(s)
    stack: List[tuple] = []
    skip = False
    uc = 1
    intbl = False
    first = star = False
    uc_pending = 0
    hexbuf = bytearray()
    hi_sur: Optional[int] = None
    pict_depth: Optional[int] = None
    pict_ok = pict_bin = False
    pict_hex: List[str] = []

    def flush():
        nonlocal hexbuf
        out = []
        if hexbuf:
            out.append(bytes(hexbuf).decode(enc, "replace"))
            hexbuf = bytearray()
        return out

    def text_out(t: str):
        """Текст с учётом недописанного суррогата."""
        nonlocal hi_sur
        res = flush()
        if hi_sur is not None:
            res.append("\ufffd")
            hi_sur = None
        if t:
            res.append(t)
        return [("text", x) for x in res if x]

    for hexv, word, num, sym, brace, text in _iter_tokens(s):
        if hexv is not None:
            if skip:
                continue
            if uc_pending:
                uc_pending -= 1
                continue
            if hi_sur is not None:
                yield from text_out("")
            hexbuf.append(int(hexv, 16))
            continue
        # всё, что не \'xx, завершает последовательность байтов
        if hexbuf:
            for x in flush():
                yield ("text", x)
        if brace == "{":
            stack.append((skip, uc, intbl))
            first, star = True, False
            uc_pending = 0
            continue
        if brace == "}":
            uc_pending = 0
            if pict_depth is not None and len(stack) == pict_depth:
                if pict_ok and not pict_bin:
                    data = _HEX_DIGITS.sub("", "".join(pict_hex))
                    if data:
                        data = data[: len(data) // 2 * 2]
                        yield ("img", hashlib.sha1(bytes.fromhex(data)).hexdigest()[:12])
                pict_depth = None
                pict_hex = []
            skip, uc, intbl = stack.pop() if stack else (False, 1, False)
            first = star = False
            continue
        if sym is not None:
            if sym == "*" and first:
                skip, star = True, True
                continue
            first = False
            if skip:
                continue
            if uc_pending:
                uc_pending -= 1
                continue
            if sym in "{}\\":
                yield from text_out(sym)
            elif sym == "~":
                yield from text_out(" ")
            elif sym == "_":
                yield from text_out("-")
            continue
        if word is not None:
            if first:
                if star and word == "shppict":
                    skip = stack[-1][0] if stack else False      # современная картинка Word
                elif word in _SKIP_DEST:
                    skip = True
                    if word == "pict":
                        pict_depth = len(stack)
                        pict_ok = not (stack[-1][0] if stack else False)   # родитель не пропускается
                        pict_bin, pict_hex = False, []
            first = star = False
            if word == "bin":
                if pict_depth is not None and len(stack) == pict_depth:
                    if pict_ok:
                        digest = hashlib.sha1(num.encode(enc, "surrogateescape")).hexdigest()[:12]
                        yield ("img", digest)
                    pict_bin = True
                continue
            if skip:
                continue
            if word == "u" and num is not None:
                code = int(num)
                if code < 0:
                    code += 65536
                if 0xD800 <= code <= 0xDBFF:
                    if hi_sur is not None:
                        yield from text_out("")
                    hi_sur = code
                elif 0xDC00 <= code <= 0xDFFF:
                    if hi_sur is not None:
                        ch = chr(0x10000 + ((hi_sur - 0xD800) << 10) + (code - 0xDC00))
                        hi_sur = None
                        yield ("text", ch)
                    else:
                        yield from text_out("\ufffd")
                else:
                    yield from text_out(chr(code))
                uc_pending = uc
                continue
            uc_pending = 0
            if word == "uc":
                uc = int(num) if num is not None else 1
            elif word == "pard":
                intbl = False
            elif word == "intbl":
                intbl = True
            elif word == "itap":
                intbl = int(num or 1) > 0
            elif word in _STRUCT:
                if hi_sur is not None:
                    yield from text_out("")
                yield (word, intbl) if word in ("par", "trowd") else (word,)
            elif word in _SYMBOLS:
                yield from text_out(_SYMBOLS[word])
            continue
        # обычный текст
        first = star = False
        if pict_depth is not None and len(stack) == pict_depth and not pict_bin:
            pict_hex.append(text)
        if skip:
            continue
        t = text.replace("\r", "").replace("\n", "")
        if uc_pending and t:
            k = min(uc_pending, len(t))
            t, uc_pending = t[k:], uc_pending - k
        if t:
            yield from text_out(t)
    yield from text_out("")
    yield ("end", intbl)


# ───────────────────────── картинки вместо текста ─────────────────────────
# Справочные системы выносят в картинки (объекты MathType) иероглифы, якутские и
# латинские слова с диакритикой. Тексты картинок хранятся в image_text.json
# (ключ — первые 12 знаков SHA-1 содержимого картинки).

IMG_RE = re.compile(r"⟦img:([0-9a-f]{12})⟧")
UNRESOLVED: Dict[str, str] = {}      # хэш → где встретилась (для предупреждений)


def _load_image_text() -> Dict[str, str]:
    import json
    p = Path(__file__).resolve().parent / "image_text.json"
    if not p.exists():
        return {}
    data = json.loads(p.read_text(encoding="utf-8"))
    return {k: v for k, v in data.items() if not k.startswith("_")}


IMAGE_TEXT = _load_image_text()


def resolve_images(s: str) -> str:
    def sub(m):
        t = IMAGE_TEXT.get(m.group(1))
        if t is None:
            UNRESOLVED[m.group(1)] = s[:80]
            return "⟦рисунок⟧"
        return t
    out = IMG_RE.sub(sub, s)
    return re.sub(r"\s*\x08", "", out)       # \b в тексте картинки: «приклеить к предыдущему»


def rtf_rows_raw(s: str) -> List[List[str]]:
    """
    Строки таблиц RTF как списки сырых текстов ячеек (картинки — ⟦img:sha⟧).
      • абзацы ячейки — через \\n; ячейки вложенной таблицы (\\nestcell) — тоже через \\n;
      • буфер ячейки сбрасывается на \\row, в конце абзаца вне таблицы (\\pard без
        \\intbl) и на \\trowd вне таблицы, поэтому текст между таблицами не прилипает
        к первой ячейке;
      • незавершённая строка (нет \\row) сохраняется, если в ней есть текст.
    """
    rows: List[List[str]] = []
    cells: List[str] = []
    buf: List[str] = []
    started = False

    def close_row():
        nonlocal cells
        if cells and any(c.strip() for c in cells):
            rows.append(cells)
        cells = []

    for ev in _rtf_events(s):
        kind = ev[0]
        if kind == "text":
            if started:
                buf.append(ev[1])
        elif kind == "img":
            if started:
                buf.append(f"⟦img:{ev[1]}⟧")
        elif kind == "trowd":
            started = True
            if not ev[1] and not cells:    # текст вне таблицы перед новой строкой таблицы
                buf = []
        elif kind == "par":
            if ev[1]:
                buf.append("\n")
            else:                      # абзац вне таблицы закончился
                close_row()
                buf = []
        elif kind in ("line", "nestcell", "nestrow"):
            buf.append("\n")
        elif kind == "tab":
            buf.append(" ")
        elif kind == "cell":
            started = True
            cells.append("".join(buf))
            buf = []
        elif kind == "row":
            rows.append(cells)
            cells, buf = [], []
        elif kind == "end":
            tail = "".join(buf)
            if ev[1] and tail.strip():
                cells.append(tail)
            close_row()
    return rows


def rtf_tables(s: str) -> List[List[str]]:
    """Все строки всех таблиц RTF как списки текстов ячеек (абзацы ячейки — через \\n)."""
    return [[clean_cell(resolve_images(c)) for c in r] for r in rtf_rows_raw(s)]


def rtf_plain(s: str, limit: int = 60000) -> str:
    """Весь видимый текст RTF (вне зависимости от таблиц) — для реквизитов приказа."""
    out: List[str] = []
    size = 0
    for ev in _rtf_events(s):
        kind = ev[0]
        if kind == "text":
            out.append(ev[1])
            size += len(ev[1])
            if size > limit:
                break
        elif kind in ("par", "line", "row", "nestrow"):
            out.append("\n")
        elif kind in ("cell", "tab", "nestcell"):
            out.append(" ")
    return "".join(out)


# ───────────────────────── DOCX ─────────────────────────

def docx_tables(path: Path):
    from docx import Document
    d = Document(str(path))
    rows = []
    for t in d.tables:
        for r in t.rows:
            rows.append([clean_cell(c.text) for c in r.cells])
    head = "\n".join(p.text for p in d.paragraphs if p.text.strip())
    return rows, head


# ───────────────────────── реквизиты ─────────────────────────

MINISTRIES = {
    "Минобрнауки": "Минобрнауки России",
    "Минпросвещения": "Минпросвещения России",
}


def parse_meta(text: str) -> Dict[str, str]:
    """Реквизиты приказа из текста заголовка («Приказ Минобрнауки России от … N …»)."""
    t = text.replace("\xa0", " ")
    meta: Dict[str, str] = {}
    m = re.search(r"Приказ\s+(Мин[а-я]+)\s+России\s+от\s+(\d\d\.\d\d\.\d{4})\s+N\s*([\w/-]+)", t)
    if m:
        meta["ministry"] = MINISTRIES.get(m.group(1), m.group(1) + " России")
        meta["date"], meta["number"] = m.group(2), m.group(3)
    m = re.search(r'"\s*(Об утверждении.*?)\s*"\s*\(Зарегистрировано', t, re.S)
    if m:
        meta["title"] = re.sub(r"\s+", " ", m.group(1)).strip()
    m = re.search(r"Зарегистрировано в Минюсте России\s+(\d\d\.\d\d\.\d{4})\s+N\s*(\d+)", t)
    if m:
        meta["reg_date"], meta["reg_number"] = m.group(1), m.group(2)
    m = re.search(r"Начало действия документа\s*-\s*(\d\d\.\d\d\.\d{4})", t)
    if m:
        meta["effective_date"] = m.group(1)
    return meta


def format_source(meta: Dict[str, str]) -> str:
    """Строка-ссылка на приказ для карточки и справки."""
    if not meta.get("number"):
        return ""
    s = f"Приказ {meta.get('ministry', '')} от {meta['date']} № {meta['number']}"
    if meta.get("title"):
        s += f" «{meta['title']}»"
    extra = []
    if meta.get("reg_number"):
        extra.append(f"зарегистрирован в Минюсте России {meta['reg_date']} № {meta['reg_number']}")
    if meta.get("effective_date"):
        extra.append(f"действует с {meta['effective_date']}")
    if extra:
        s += " (" + "; ".join(extra) + ")"
    return s


# ───────────────────────── точка входа ─────────────────────────

@dataclass
class Order:
    path: str
    meta: Dict[str, str] = field(default_factory=dict)
    rows: List[List[str]] = field(default_factory=list)
    # картинки без текста в image_text.json: (номер строки в rows с 1, № в приказе, sha1[:12])
    unresolved_images: List[Tuple[int, str, str]] = field(default_factory=list)


def _keep_row(r: List[str]) -> bool:
    """Строки таблицы с данными (≥ 6 ячеек) и заголовки разделов (слитая ячейка)."""
    return len(r) >= 6 or (bool(r) and is_section_row(r) is not None)


def read_order(path) -> Order:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(path)
    suffix = p.suffix.lower()
    unresolved: List[Tuple[int, str, str]] = []
    if suffix == ".rtf":
        s = _decode_rtf(p.read_bytes())
        rows = []
        for raw in rtf_rows_raw(s):
            r = [clean_cell(resolve_images(c)) for c in raw]
            if not _keep_row(r):
                continue
            rows.append(r)
            if len(r) >= 6:
                for sha in uniq(m for c in raw for m in IMG_RE.findall(c)):
                    if sha not in IMAGE_TEXT:
                        unresolved.append((len(rows), r[0], sha))
        meta_text = rtf_plain(s)
    elif suffix == ".docx":
        rows, head = docx_tables(p)
        rows = [r for r in rows if _keep_row(r)]
        meta_text = head
    else:
        raise ValueError(f"Неподдерживаемый формат: {p.suffix} (нужен .rtf или .docx)")
    return Order(path=str(p), meta=parse_meta(meta_text), rows=rows, unresolved_images=unresolved)


def image_warnings(order: Order) -> List[str]:
    """Предупреждения о картинках в строках таблицы, текста которых нет в image_text.json."""
    return [f"картинка без расшифровки в строке {n}{f' (№ {num})' if num else ''} "
            f"(sha1 {sha}) — добавьте в image_text.json"
            for n, num, sha in order.unresolved_images]


def find_order(ministry: str, folder: str = ".") -> Optional[str]:
    """Находит в папке приказ нужного министерства (.rtf предпочтительнее .docx)."""
    found = []
    for p in sorted(Path(folder).iterdir()):
        if p.suffix.lower() not in (".rtf", ".docx") or p.name.startswith("~$"):
            continue
        try:
            o = read_order(p)
        except Exception:  # noqa: BLE001
            continue
        if o.meta.get("ministry") == ministry and o.meta.get("number"):
            found.append((p.suffix.lower() == ".rtf", p.stat().st_mtime, str(p)))
    return max(found)[2] if found else None


# ───────────────────────── общие приёмы разбора строк ─────────────────────────

def pad_row(r: List[str], n: int = 6) -> List[str]:
    """Строка, дополненная пустыми ячейками до n (заголовки разделов в RTF — одна ячейка)."""
    return list(r) + [""] * (n - len(r)) if len(r) < n else list(r)


def is_header_row(r: List[str]) -> bool:
    r = pad_row(r)
    n, name = r[0], r[1]
    if ("п/п" in n) or n.lower().startswith(("n п", "№")):
        return True
    if name.startswith(("Полное наименование", "Название мероприятия", "Наименование")):
        return True
    if r[3].startswith(("Профиль олимпиады", "Профильное направление", "Направление мероприятия")):
        return True
    if r[4].startswith(("Общеобразовательные предметы", "Профильное направление")):
        return True
    return False


SECTION_RE = re.compile(r"^(\d+)\.\s+Мероприятия\b")


def is_section_row(r: List[str]) -> Optional[int]:
    """
    Номер раздела, если это строка-заголовок раздела: в RTF — строка из одной слитой
    ячейки, в docx — текст повторён во всех ячейках.
    """
    if not r:
        return None
    m = SECTION_RE.match(r[0])
    second = r[1] if len(r) > 1 else ""
    if m and (not second or second == r[0]) and not any(c and c != r[0] for c in r[2:]):
        return int(m.group(1))
    return None


LEVEL_RX = [
    (re.compile(r"^\s*высш", re.I), "Высший"),
    (re.compile(r"^\s*(IV|III|II|I)\b", re.I), None),      # римская цифра как есть
    (re.compile(r"^\s*не\s+установлен", re.I), "Не установлен"),
    (re.compile(r"^\s*не\s*\.?\s*$", re.I), "Не установлен"),   # в приказе № 598 (6.345) слово обрезано до «Не»
]


def norm_level(s: str) -> str:
    """«IV уровень» → «IV», «Высший уровень» → «Высший», «Не установлен.» → «Не установлен»."""
    s = (s or "").strip()
    for rx, val in LEVEL_RX:
        m = rx.match(s)
        if m:
            return val or m.group(1).upper()
    return s


def split_top_level(s: str, sep: str = ";") -> List[str]:
    """Делит по разделителю вне скобок и кавычек; пустые части отбрасывает."""
    parts, buf, depth, quote = [], [], 0, 0
    for ch in s:
        if ch in "([":
            depth += 1
        elif ch in ")]":
            depth = max(0, depth - 1)
        elif ch == "«":
            quote += 1
        elif ch == "»":
            quote = max(0, quote - 1)
        if ch == sep and depth == 0 and quote == 0:
            parts.append("".join(buf).strip())
            buf = []
        else:
            buf.append(ch)
    parts.append("".join(buf).strip())
    return [p for p in parts if p]


def uniq(seq):
    return list(dict.fromkeys(x for x in seq if x))
